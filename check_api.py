"""Exercise the real audio API and retain responses and checks in results/.

Each endpoint performs its own speech recognition and meeting analysis. Transcript
comparisons therefore report differences between independent inference runs; they
do not prove that a difference originated in the document exporter.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
import hashlib
from io import BytesIO
import json
import math
import mimetypes
from pathlib import Path
import re
import time
import unicodedata
from urllib.parse import urlsplit
from zipfile import ZipFile


ROUTES = {
    "json": "/process-audio",
    "pdf": "/process-audio-pdf",
    "docx": "/process-audio-docx",
}
MIME_TYPES = {
    "json": "application/json",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
SAFE_HEADERS = ("Content-Type", "Content-Disposition", "Cache-Control", "Retry-After")
REPORT_TITLE = "Поручения по итогам совещания"
TRANSCRIPT_HEADING = "Расшифровка совещания"


def timestamp() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: object) -> None:
    # Replace the complete manifest, so an interrupted subsequent request leaves
    # the checks for completed routes readable.
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def normalize_text(value: str) -> str:
    # PDF wraps paragraphs and may split a long word across lines. Preserve all
    # letters and punctuation while ignoring only layout whitespace/soft hyphens.
    return "".join(unicodedata.normalize("NFC", value).replace("\u00ad", "").split())


def check(entry: dict, name: str, passed: bool, detail: object = None) -> None:
    item = {"name": name, "passed": bool(passed)}
    if detail is not None:
        item["detail"] = detail
    entry["checks"].append(item)


def validate_json(data: bytes, entry: dict, args: argparse.Namespace) -> str | None:
    report = json.loads(data)
    check(entry, "report_is_object", isinstance(report, dict))
    if not isinstance(report, dict):
        return None
    transcript = report.get("transcript")
    check(entry, "transcript_is_object", isinstance(transcript, dict))
    if not isinstance(transcript, dict):
        return None
    text = transcript.get("text")
    nonempty = isinstance(text, str) and bool(text.strip())
    check(entry, "transcript_nonempty", nonempty)
    duration = transcript.get("duration")
    check(entry, "duration_positive", type(duration) in (int, float) and math.isfinite(duration) and duration > 0,
          duration)
    check(entry, "filename_matches", report.get("filename") == args.audio.name, report.get("filename"))
    check(entry, "meeting_date_matches", report.get("meeting_date") == args.meeting_date, report.get("meeting_date"))
    check(entry, "analysis_method_present", isinstance(report.get("analysis_method"), str)
          and bool(report["analysis_method"].strip()), report.get("analysis_method"))
    tasks, dashboard = report.get("tasks"), report.get("dashboard")
    valid_tasks = isinstance(tasks, list) and all(isinstance(task, dict) for task in tasks)
    check(entry, "tasks_is_object_list", valid_tasks)
    check(entry, "dashboard_is_object", isinstance(dashboard, dict))
    if valid_tasks and isinstance(dashboard, dict):
        check(entry, "dashboard_total", type(dashboard.get("total")) is int and dashboard["total"] == len(tasks))
        for key, allowed in (("status", {"in_progress", "overdue", "completed"}),
                             ("urgency", {"high", "medium", "low"})):
            check(entry, f"task_{key}_values", all(task.get(key) in allowed for task in tasks))
            expected = {value: sum(task.get(key) == value for task in tasks) for value in sorted(allowed)}
            actual = dashboard.get(f"by_{key}")
            check(entry, f"dashboard_by_{key}", actual == expected
                  and isinstance(actual, dict) and all(type(value) is int for value in actual.values()),
                  {"expected": expected, "actual": actual})
        valid_directions = all(isinstance(task.get("direction"), str) and task["direction"] for task in tasks)
        check(entry, "task_direction_values", valid_directions)
        if valid_directions:
            expected = dict(Counter(task["direction"] for task in tasks))
            actual = dashboard.get("by_direction")
            check(entry, "dashboard_by_direction", actual == expected
                  and isinstance(actual, dict) and all(type(value) is int for value in actual.values()),
                  {"expected": expected, "actual": actual})
        check(entry, "task_needs_review_boolean", all(type(task.get("needs_review")) is bool for task in tasks))
        expected_review = sum(task.get("needs_review") is True for task in tasks)
        check(entry, "dashboard_needs_review", type(dashboard.get("needs_review")) is int
              and dashboard["needs_review"] == expected_review,
              {"expected": expected_review, "actual": dashboard.get("needs_review")})
        entry["task_count"] = len(tasks)
    if args.diarize == "true" and nonempty:
        speakers, turns = transcript.get("speakers"), transcript.get("speaker_turns")
        valid_speakers = isinstance(speakers, list) and bool(speakers) and all(
            isinstance(speaker, str) and bool(speaker.strip()) for speaker in speakers)
        check(entry, "diarization_speakers_present", valid_speakers, speakers)
        valid_turns = isinstance(turns, list) and bool(turns) and all(isinstance(turn, dict) for turn in turns)
        check(entry, "diarization_turns_present", valid_turns)
        if valid_speakers and valid_turns:
            check(entry, "diarization_turn_speakers_known", all(turn.get("speaker") in speakers for turn in turns))
        entry["speaker_count"] = len(speakers) if isinstance(speakers, list) else None
    if nonempty:
        (args.output_dir / "transcript.txt").write_text(text + "\n", encoding="utf-8")
        entry["transcript_characters"] = len(text)
        return text
    return None


def compare_transcript(extracted: str, reference: str | None, entry: dict) -> None:
    normalized = normalize_text(extracted)
    check(entry, "document_transcript_nonempty", bool(normalized)
          and normalized != normalize_text("Расшифровка отсутствует."))
    entry["transcript_characters"] = len(extracted)
    if reference is None:
        entry["transcript_comparison"] = {
            "available": False,
            "reason": "No matching successful JSON response is available; run the json route first.",
        }
        return
    expected = normalize_text(reference)
    matches = normalized == expected
    comparison = {"available": True, "normalized_full_text_equal": matches,
                  "expected_characters": len(expected), "actual_characters": len(normalized),
                  "expected_sha256": sha256(expected.encode("utf-8")),
                  "actual_sha256": sha256(normalized.encode("utf-8"))}
    if not matches:
        comparison["similarity"] = round(SequenceMatcher(None, expected, normalized).ratio(), 6)
        comparison["note"] = ("Endpoints run inference independently. A mismatch needs review and does not "
                              "by itself establish that the exporter lost text.")
    entry["transcript_comparison"] = comparison
    check(entry, "full_transcript_matches_json", matches, comparison)


def validate_document(kind: str, data: bytes, entry: dict, args: argparse.Namespace,
                      reference: str | None) -> None:
    if kind == "pdf":
        check(entry, "pdf_signature", data.startswith(b"%PDF-"))
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        check(entry, "pdf_has_pages", len(reader.pages) > 0)
        check(entry, "pdf_not_encrypted", not reader.is_encrypted)
        entry["page_count"] = len(reader.pages)
        # Footers are document layout, not recognized speech. Remove one exact
        # footer line per page, leaving any identical line in the transcript.
        pages = []
        for number, page in enumerate(reader.pages, 1):
            page_text = page.extract_text() or ""
            page_text = re.sub(rf"(?m)^Страница\s+{number}\s*$", "", page_text, count=1)
            pages.append(page_text)
        text = "\n".join(pages)
    else:
        check(entry, "docx_zip_signature", data.startswith(b"PK\x03\x04"))
        with ZipFile(BytesIO(data)) as archive:
            check(entry, "docx_zip_integrity", archive.testzip() is None)
            names = set(archive.namelist())
            check(entry, "docx_required_parts", {"[Content_Types].xml", "word/document.xml"}.issubset(names))
        from docx import Document
        document = Document(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        entry["paragraph_count"] = len(document.paragraphs)
    (args.output_dir / f"meeting-report-{kind}.txt").write_text(text + "\n", encoding="utf-8")
    check(entry, "report_title_present", normalize_text(REPORT_TITLE) in normalize_text(text))
    check(entry, "transcript_heading_present", TRANSCRIPT_HEADING in text)
    check(entry, "meeting_date_present", args.meeting_date in text)
    if TRANSCRIPT_HEADING in text:
        # The transcript is the final section. Do not accidentally count a task's
        # supporting source quote as proof that the full transcript was exported.
        extracted = text.split(TRANSCRIPT_HEADING, 1)[1].strip()
        compare_transcript(extracted, reference, entry)


def previous_reference(args: argparse.Namespace, audio_digest: str) -> str | None:
    if "json" in args.routes:
        return None
    try:
        manifest = json.loads((args.output_dir / "validation.json").read_text(encoding="utf-8"))
        report_path = args.output_dir / "meeting-report.json"
        data = report_path.read_bytes()
        json_result = manifest["routes"]["json"]
        if (manifest["audio"]["sha256"] != audio_digest
                or manifest["base_url"] != args.base_url
                or manifest["meeting_date"] != args.meeting_date
                or manifest["diarize"] != args.diarize
                or json_result["sha256"] != sha256(data)
                or not json_result["passed"]):
            return None
        text = json.loads(data)["transcript"]["text"]
        return text if isinstance(text, str) and text.strip() else None
    except (OSError, ValueError, KeyError, TypeError):
        return None


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--audio", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=Path("results/api-check"))
    parser.add_argument("--meeting-date", default=date.today().isoformat(), help="YYYY-MM-DD (default: today)")
    parser.add_argument("--diarize", choices=("true", "false"), default="true")
    parser.add_argument("--routes", choices=tuple(ROUTES), nargs="+", default=list(ROUTES))
    args = parser.parse_args(argv)
    if not args.audio.is_file() or args.audio.stat().st_size == 0:
        parser.error("--audio must be an existing, nonempty recording")
    try:
        if date.fromisoformat(args.meeting_date).isoformat() != args.meeting_date:
            raise ValueError
    except ValueError:
        parser.error("--meeting-date must be a valid YYYY-MM-DD date")
    parsed_url = urlsplit(args.base_url)
    if (parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc
            or parsed_url.username or parsed_url.password or parsed_url.query or parsed_url.fragment):
        parser.error("--base-url must be an HTTP(S) API URL without credentials, query or fragment")
    args.base_url = args.base_url.rstrip("/")
    args.routes = [kind for kind in ROUTES if kind in args.routes]
    return args


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    # Fail before expensive inference if a requested format cannot be inspected.
    import requests
    if "pdf" in args.routes:
        from pypdf import PdfReader  # noqa: F401
    if "docx" in args.routes:
        from docx import Document  # noqa: F401

    args.output_dir.mkdir(parents=True, exist_ok=True)
    audio_digest = file_sha256(args.audio)
    reference = previous_reference(args, audio_digest)
    retained = {}
    if reference is not None:
        previous = json.loads((args.output_dir / "validation.json").read_text(encoding="utf-8"))
        for kind, result in previous["routes"].items():
            if kind not in ROUTES or kind in args.routes or not result.get("passed"):
                continue
            path = args.output_dir / f"meeting-report.{kind}"
            if path.is_file() and file_sha256(path) == result.get("sha256"):
                retained[kind] = result
    manifest = {"started_at": timestamp(), "base_url": args.base_url,
                "audio": {"path": str(args.audio.resolve()), "bytes": args.audio.stat().st_size,
                          "sha256": audio_digest},
                "meeting_date": args.meeting_date, "diarize": args.diarize,
                "requested_routes": args.routes, "routes": retained, "passed": False,
                "note": "Real multipart HTTP requests; each route runs its own inference."}
    manifest_path = args.output_dir / "validation.json"
    write_json(manifest_path, manifest)
    with requests.Session() as session:
        session.trust_env = False
        for kind in args.routes:
            route = ROUTES[kind]
            output = args.output_dir / f"meeting-report.{kind}"
            error_path = args.output_dir / f"{kind}-error.json"
            # Remove only this checker's known stale output for a rerun.
            output.unlink(missing_ok=True)
            error_path.unlink(missing_ok=True)
            if kind == "json":
                (args.output_dir / "transcript.txt").unlink(missing_ok=True)
            else:
                (args.output_dir / f"meeting-report-{kind}.txt").unlink(missing_ok=True)
            print(f"POST {args.base_url}{route}: {args.audio.name} (diarize={args.diarize})", flush=True)
            entry = {"route": route, "started_at": timestamp(), "status": None, "headers": {},
                     "elapsed_seconds": None, "bytes": 0, "sha256": None, "checks": [], "passed": False}
            started = time.monotonic()
            response = None
            try:
                with args.audio.open("rb") as audio:
                    response = session.post(
                        args.base_url + route,
                        files={"audio": (args.audio.name, audio, mimetypes.guess_type(args.audio.name)[0]
                                         or "application/octet-stream")},
                        data={"meeting_date": args.meeting_date, "diarize": args.diarize},
                        timeout=(10, 1800), allow_redirects=False,
                    )
                data = response.content
                entry.update(status=response.status_code,
                             headers={key: response.headers[key] for key in SAFE_HEADERS if key in response.headers},
                             elapsed_seconds=round(time.monotonic() - started, 3),
                             bytes=len(data), sha256=sha256(data))
                check(entry, "http_status_200", response.status_code == 200)
                if response.status_code != 200:
                    raise ValueError(f"HTTP {response.status_code}")
                output.write_bytes(data)
                entry["output"] = output.name
                content_type = response.headers.get("Content-Type", "").split(";", 1)[0].strip().lower()
                check(entry, "content_type", content_type == MIME_TYPES[kind], content_type)
                if kind == "json":
                    reference = validate_json(data, entry, args)
                else:
                    disposition = response.headers.get("Content-Disposition", "")
                    check(entry, "download_attachment", "attachment" in disposition.lower())
                    validate_document(kind, data, entry, args, reference)
                entry["passed"] = all(item["passed"] for item in entry["checks"])
                if not entry["passed"]:
                    raise ValueError("Failed checks: " + ", ".join(
                        item["name"] for item in entry["checks"] if not item["passed"]))
            except Exception as exc:
                entry["error"] = f"{type(exc).__name__}: {exc}"
                error = {"route": route, "error": entry["error"], "status": entry["status"]}
                if response is not None and response.status_code != 200:
                    try:
                        error["response"] = response.json()
                    except ValueError:
                        error["response_text"] = response.text
                error["checks"] = entry["checks"]
                write_json(error_path, error)
                entry["error_output"] = error_path.name
            finally:
                if entry["elapsed_seconds"] is None:
                    entry["elapsed_seconds"] = round(time.monotonic() - started, 3)
                entry["finished_at"] = timestamp()
                manifest["routes"][kind] = entry
                manifest["updated_at"] = timestamp()
                manifest["passed"] = all(kind in manifest["routes"] for kind in args.routes) and all(
                    result["passed"] for result in manifest["routes"].values())
                write_json(manifest_path, manifest)
            print(f"{kind}: {'PASS' if entry['passed'] else 'FAIL'}; HTTP {entry['status']}; "
                  f"{entry['elapsed_seconds']}s; {entry['bytes']} bytes", flush=True)
            if entry.get("error"):
                print(entry["error"], flush=True)
    print(f"Validation: {manifest_path.resolve()}", flush=True)
    return 0 if manifest["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
