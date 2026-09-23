"""Render one meeting analysis as a PDF or DOCX, entirely in memory.

The API owns analysis and status calculation. These exporters consume the same
result dictionary as the JSON endpoint so downloads never change task semantics.
"""

from __future__ import annotations

import hashlib
import os
import re
import threading
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape


STATUS_LABELS = {
    "in_progress": "В работе", "overdue": "Просрочено", "completed": "Выполнено"
}
URGENCY_LABELS = {"high": "Высокая", "medium": "Средняя", "low": "Низкая"}
DIRECTION_LABELS = {
    "finance": "Финансы", "legal": "Право", "procurement": "Закупки",
    "production": "Производство", "safety": "Безопасность", "hr": "Персонал",
    "it": "ИТ", "general": "Общие вопросы",
}
_FONT_LOCK = threading.Lock()
_PDF_RENDER_LOCK = threading.Lock()
_REQUIRED_GLYPHS = "АБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯӘҒҚҢӨҰҮҺІабвгдеёжзийклмнопрстуфхцчшщъыьэюяәғқңөұүһі"
_INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


class ReportFontError(RuntimeError):
    """The configured report fonts are absent or cannot render RU/KZ text."""


@dataclass(frozen=True)
class _Block:
    kind: str
    text: str = ""
    label: str = ""


def _clean(value) -> str:
    return _INVALID_XML.sub("", str(value)).replace("\r\n", "\n").replace("\r", "\n")


def _value(value, default="Не указано") -> str:
    return _clean(value) if value is not None and str(value).strip() else default


def _font_paths() -> tuple[Path, Path]:
    configured = os.environ.get("REPORT_FONT_PATH")
    bold_configured = os.environ.get("REPORT_FONT_BOLD_PATH")
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [
        (windows_fonts / "arial.ttf", windows_fonts / "arialbd.ttf"),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"), Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf")),
        (Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"), Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")),
        (Path("/Library/Fonts/Arial.ttf"), Path("/Library/Fonts/Arial Bold.ttf")),
        (Path("/System/Library/Fonts/Supplemental/Arial.ttf"), Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")),
    ]
    if configured:
        regular = Path(configured).expanduser()
        if not regular.is_file():
            raise ReportFontError("REPORT_FONT_PATH must point to an existing Unicode TrueType font.")
        bold = regular
    else:
        pair = next((pair for pair in candidates if pair[0].is_file()), None)
        if pair is None:
            raise ReportFontError(
                "No Unicode report font found. Install DejaVu Sans, Liberation Sans or Arial, "
                "or set REPORT_FONT_PATH (and optionally REPORT_FONT_BOLD_PATH)."
            )
        regular, bold = pair
        if not bold.is_file():
            bold = regular
    if bold_configured:
        bold = Path(bold_configured).expanduser()
        if not bold.is_file():
            raise ReportFontError("REPORT_FONT_BOLD_PATH must point to an existing Unicode TrueType font.")
    return regular, bold


def _fonts() -> tuple[str, str, str]:
    """Register embedded PDF fonts once; also resolve the DOCX font family."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    paths = _font_paths()
    names = []
    family = ""
    # ReportLab has a process-wide font registry; serialize its initialization.
    with _FONT_LOCK:
        for index, path in enumerate(paths):
            name = "MeetingReport" + hashlib.sha256(str(path.resolve()).encode()).hexdigest()[:16]
            try:
                if name not in pdfmetrics.getRegisteredFontNames():
                    font = TTFont(name, str(path))
                    missing = [char for char in _REQUIRED_GLYPHS if ord(char) not in font.face.charToGlyph]
                    if missing:
                        raise ReportFontError(f"Report font {path.name} lacks Russian/Kazakh glyphs: {''.join(missing)}")
                    pdfmetrics.registerFont(font)
                font = pdfmetrics.getFont(name)
                if index == 0:
                    raw_family = font.face.familyName
                    family = raw_family.decode("utf-8", errors="replace") if isinstance(raw_family, bytes) else str(raw_family)
                names.append(name)
            except ReportFontError:
                raise
            except Exception as exc:
                raise ReportFontError(f"Cannot load report font {path.name}. Configure a valid TrueType font.") from exc
        pdfmetrics.registerFontFamily(names[0], normal=names[0], bold=names[1], italic=names[0], boldItalic=names[1])
    return names[0], names[1], family


def _blocks(result: dict):
    yield _Block("title", "Поручения по итогам совещания")
    yield _Block("body", "Отчёт содержит поручения, ответственных, сроки и сводку для контроля исполнения.")
    for label, key in (("Аудиозапись", "filename"), ("Дата совещания", "meeting_date"), ("Отчёт сформирован", "generated_at")):
        yield _Block("meta", _value(result.get(key)), label)
    for label, key in (("Статусы актуальны на", "as_of"), ("Часовой пояс", "timezone")):
        if result.get(key):
            yield _Block("meta", _value(result[key]), label)
    method = result.get("analysis_method")
    if method:
        yield _Block("meta", _value(method), "Метод анализа")

    dashboard = result.get("dashboard") or {}
    yield _Block("heading", "Дашборд поручений")
    yield _Block("body", str(dashboard.get("total", len(result.get("tasks") or []))), "Всего поручений")
    statuses = dashboard.get("by_status") or {}
    for code, label in STATUS_LABELS.items():
        yield _Block("body", str(statuses.get(code, 0)), label)
    yield _Block("body", str(dashboard.get("needs_review", 0)), "Требуют проверки")
    yield _Block("subheading", "По срочности")
    urgencies = dashboard.get("by_urgency") or {}
    for code, label in URGENCY_LABELS.items():
        yield _Block("body", str(urgencies.get(code, 0)), label)
    yield _Block("subheading", "По направлениям")
    directions = dashboard.get("by_direction") or {}
    if directions:
        for label, count in sorted(directions.items()):
            yield _Block("body", str(count), DIRECTION_LABELS.get(label, _clean(label)))
    else:
        yield _Block("body", "Направления не определены.")

    warnings = result.get("warnings") or []
    if warnings:
        yield _Block("heading", "Примечания к анализу")
        for warning in warnings:
            yield _Block("body", _clean(warning))

    yield _Block("heading", "Список поручений")
    tasks = result.get("tasks") or []
    if not tasks:
        yield _Block("body", "Поручения не обнаружены.")
    for index, task in enumerate(tasks, 1):
        yield _Block("subheading", f"Поручение {index}")
        yield _Block("task_title", _value(task.get("title")))
        fields = [
            ("Идентификатор", _value(task.get("id"))),
            ("Ответственный", _value(task.get("assignee"), "Не указан")),
            ("Срок", _value(task.get("deadline"), "Не указан")),
            ("Срок в записи", _value(task.get("deadline_text"), "Не указан")),
            ("Статус", STATUS_LABELS.get(task.get("status"), _value(task.get("status")))),
            ("Срочность", URGENCY_LABELS.get(task.get("urgency"), _value(task.get("urgency")))),
            ("Направление", DIRECTION_LABELS.get(task.get("direction"), _value(task.get("direction")))),
            ("Требует проверки", "Да" if task.get("needs_review") else "Нет"),
            ("Говорящий", _value(task.get("source_speaker"), "Не определён")),
            ("Основание из расшифровки", _value(task.get("source_quote"))),
        ]
        for label, value in fields:
            yield _Block("body", value, label)

    yield _Block("heading", "Расшифровка совещания")
    transcript = result.get("transcript") or {}
    text = _value(transcript.get("text") if isinstance(transcript, dict) else transcript, "Расшифровка отсутствует.")
    for paragraph in text.split("\n"):
        if paragraph.strip():
            yield _Block("body", paragraph)


def render_pdf(result: dict) -> bytes:
    """Create a paginated PDF with embedded Russian/Kazakh TrueType fonts."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    regular, bold, _ = _fonts()
    styles = {
        "body": ParagraphStyle("body", fontName=regular, fontSize=11, leading=15, spaceAfter=5, alignment=TA_LEFT, splitLongWords=True),
        "meta": ParagraphStyle("meta", fontName=regular, fontSize=9, leading=13, spaceAfter=3, textColor=colors.HexColor("#475569")),
        "title": ParagraphStyle("title", fontName=bold, fontSize=21, leading=26, spaceAfter=14, keepWithNext=True),
        "heading": ParagraphStyle("heading", fontName=bold, fontSize=15, leading=20, spaceBefore=17, spaceAfter=8, keepWithNext=True),
        "subheading": ParagraphStyle("subheading", fontName=bold, fontSize=12, leading=16, spaceBefore=11, spaceAfter=6, keepWithNext=True),
        "task_title": ParagraphStyle("task_title", fontName=bold, fontSize=11, leading=15, spaceAfter=7, splitLongWords=True),
    }
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                                 topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                                 title="Поручения по итогам совещания", author="", pageCompression=1)
    story = []
    for block in _blocks(result):
        text = escape(_clean(block.text)).replace("\n", "<br/>")
        if block.label:
            text = f"<b>{escape(_clean(block.label))}:</b> {text}"
        story.append(Paragraph(text, styles[block.kind]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 9)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(letter[0] - 0.8 * inch, 0.35 * inch, f"Страница {doc.page}")
        canvas.restoreState()

    # ReportLab TTFont subsetting mutates the registered font's parser state.
    # Its shared font objects cannot be used by concurrent document builds.
    with _PDF_RENDER_LOCK:
        document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def render_docx(result: dict) -> bytes:
    """Create an editable DOCX with the same content as the JSON/PDF report."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    _, _, family = _fonts()
    document = Document()
    document.core_properties.title = "Поручения по итогам совещания"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name, size in (("Normal", 11), ("Title", 21), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = family
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.widow_control = True
        if name != "Normal":
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 9)
    for block in _blocks(result):
        style = {"title": "Title", "heading": "Heading 1", "subheading": "Heading 2"}.get(block.kind, "Normal")
        paragraph = document.add_paragraph(style=style)
        if block.label:
            paragraph.add_run(_clean(block.label) + ": ").bold = True
        run = paragraph.add_run(_clean(block.text))
        if block.kind == "task_title":
            run.bold = True
        elif block.kind == "meta":
            for meta_run in paragraph.runs:
                meta_run.font.size = Pt(9)
                meta_run.font.color.rgb = RGBColor.from_string("475569")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Страница ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
